"""
DTS Purpose of Visit & Top Destinations/Districts Parser (2018–2025)
Extracts:
1. State-level purpose of visit shares for both Visitors and Overnight Tourists (Jadual 8 in state DTS)
2. Top 5 tourist destinations per state (Jadual 8A in national DTS)
3. Top 5 administrative districts per state (Jadual 8B in national DTS)

Saves clean analytical tables to data/processed/ and data/processed/tourism_data.duckdb
"""

import glob
import os
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import duckdb
import openpyxl
import pandas as pd

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

DTS_BASE_DIR = ROOT_DIR / "data/dts"
PROCESSED_DIR = ROOT_DIR / "data/processed"
DUCKDB_PATH = PROCESSED_DIR / "tourism_data.duckdb"

STATE_METADATA: Dict[str, Dict] = {
    "JOHOR": {"name": "Johor", "code": "MY-01", "region": "Southern"},
    "KEDAH": {"name": "Kedah", "code": "MY-02", "region": "Northern"},
    "KELANTAN": {"name": "Kelantan", "code": "MY-03", "region": "East Coast"},
    "MELAKA": {"name": "Melaka", "code": "MY-04", "region": "Southern"},
    "NEGERI SEMBILAN": {"name": "Negeri Sembilan", "code": "MY-05", "region": "Central"},
    "PAHANG": {"name": "Pahang", "code": "MY-06", "region": "East Coast"},
    "PULAU PINANG": {"name": "Pulau Pinang", "code": "MY-07", "region": "Northern"},
    "PERAK": {"name": "Perak", "code": "MY-08", "region": "Northern"},
    "PERLIS": {"name": "Perlis", "code": "MY-09", "region": "Northern"},
    "SELANGOR": {"name": "Selangor", "code": "MY-10", "region": "Central"},
    "TERENGGANU": {"name": "Terengganu", "code": "MY-11", "region": "East Coast"},
    "SABAH": {"name": "Sabah", "code": "MY-12", "region": "East Malaysia"},
    "SARAWAK": {"name": "Sarawak", "code": "MY-13", "region": "East Malaysia"},
    "W.P. KUALA LUMPUR": {"name": "W.P. Kuala Lumpur", "code": "MY-14", "region": "Central"},
    "W.P. LABUAN": {"name": "W.P. Labuan", "code": "MY-15", "region": "East Malaysia"},
    "W.P. PUTRAJAYA": {"name": "W.P. Putrajaya", "code": "MY-16", "region": "Central"},
}


def match_state_from_name(raw_str: str) -> Optional[str]:
    cleaned = raw_str.strip().upper().replace("_", " ").replace(".", "")
    if "PULAU PINANG" in cleaned or "PENANG" in cleaned or "P PINANG" in cleaned:
        return "PULAU PINANG"
    if "KUALA LUMPUR" in cleaned or "KL" in cleaned:
        return "W.P. KUALA LUMPUR"
    if "PUTRAJAYA" in cleaned:
        return "W.P. PUTRAJAYA"
    if "LABUAN" in cleaned:
        return "W.P. LABUAN"
    if "NEGERI SEMBILAN" in cleaned or "N SEMBILAN" in cleaned:
        return "NEGERI SEMBILAN"
    for k in STATE_METADATA:
        if k.replace(".", "") in cleaned:
            return k
    return None


def normalize_purpose_category(raw_text: str) -> str:
    t = raw_text.lower()
    if any(w in t for w in ["saudara", "relatives", "vfr", "rakan"]):
        return "vfr"
    elif any(w in t for w in ["cuti", "holiday", "leisure", "lapang", "berehat", "relaxation"]):
        return "holiday_leisure"
    elif any(w in t for w in ["beli", "shopping"]):
        return "shopping"
    elif any(w in t for w in ["hiburan", "entertainment", "sukan", "sports", "acara", "event"]):
        return "entertainment_sports"
    elif any(w in t for w in ["perubatan", "medical", "rawatan", "wellness", "penjagaan"]):
        return "medical_wellness"
    elif any(w in t for w in ["ibadat", "religious", "worship"]):
        return "religious"
    elif any(w in t for w in ["rasmi", "business", "perniagaan", "pendidikan", "education", "official"]):
        return "business_education"
    elif any(w in t for w in ["insentif", "incentive", "lain", "others"]):
        return "incentive_others"
    return "other"


def parse_state_jadual_8(file_path: Path, year: int) -> Optional[Dict]:
    """Parse Jadual 8 for a single state workbook."""
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  [Error] Cannot load {file_path.name}: {e}")
        return None

    # Identify state
    file_stem = file_path.stem.replace(f"TABLE OF PUBLICATION DTS {year} ", "")
    matched_key = match_state_from_name(file_stem)
    if not matched_key:
        return None

    meta = STATE_METADATA[matched_key]

    # Find Jadual 8 sheet
    sheet_name = None
    for s in wb.sheetnames:
        if "8" in s and ("jad" in s.lower() or "table" in s.lower()):
            sheet_name = s
            break
    if not sheet_name:
        for s in wb.sheetnames:
            if s.strip() == "8":
                sheet_name = s
                break

    if not sheet_name:
        return None

    ws = wb[sheet_name]

    rows = list(ws.iter_rows(min_row=1, max_row=22, max_col=10, values_only=True))
    wb.close()

    v_col, v_share_col = None, None
    t_col, t_share_col = None, None

    # Detect header columns (0-indexed)
    for r in range(2, min(6, len(rows))):
        row_vals = rows[r]
        for c, val in enumerate(row_vals):
            if val and isinstance(val, str):
                v_lower = val.lower()
                if "pelawat" in v_lower or "visitor" in v_lower:
                    if v_col is None:
                        v_col = c
                        v_share_col = c + 1
                elif "pelancong" in v_lower or "tourist" in v_lower:
                    if t_col is None:
                        t_col = c
                        t_share_col = c + 1

    # Fallback to standard columns if not explicitly detected
    if v_col is None:
        if len(rows) > 6 and rows[6][0] is not None:
            v_col, v_share_col = 0, 1
            t_col, t_share_col = 3, 4
        else:
            v_col, v_share_col = 1, 2
            t_col, t_share_col = 4, 5

    visitor_shares: Dict[str, float] = {}
    tourist_shares: Dict[str, float] = {}

    for r in range(4, len(rows)):
        row_vals = rows[r]
        # Visitor
        if v_col is not None and v_col < len(row_vals) and v_share_col < len(row_vals):
            v_label = row_vals[v_col]
            v_val = row_vals[v_share_col]
            if v_label and isinstance(v_label, str) and "jumlah" not in v_label.lower() and "total" not in v_label.lower():
                if v_val is not None:
                    try:
                        cat = normalize_purpose_category(v_label)
                        visitor_shares[cat] = round(float(v_val), 2)
                    except (ValueError, TypeError):
                        pass

        # Tourist
        if t_col is not None and t_col < len(row_vals) and t_share_col < len(row_vals):
            t_label = row_vals[t_col]
            t_val = row_vals[t_share_col]
            if t_label and isinstance(t_label, str) and "jumlah" not in t_label.lower() and "total" not in t_label.lower():
                if t_val is not None:
                    try:
                        cat = normalize_purpose_category(t_label)
                        tourist_shares[cat] = round(float(t_val), 2)
                    except (ValueError, TypeError):
                        pass

    record = {
        "year": year,
        "state": meta["name"],
        "state_code": meta["code"],
        "region": meta["region"],
        "vfr_share_visitor": visitor_shares.get("vfr", 0.0),
        "vfr_share_tourist": tourist_shares.get("vfr", 0.0),
        "holiday_share_visitor": visitor_shares.get("holiday_leisure", 0.0),
        "holiday_share_tourist": tourist_shares.get("holiday_leisure", 0.0),
        "shopping_share_visitor": visitor_shares.get("shopping", 0.0),
        "shopping_share_tourist": tourist_shares.get("shopping", 0.0),
        "entertainment_share_visitor": visitor_shares.get("entertainment_sports", 0.0),
        "entertainment_share_tourist": tourist_shares.get("entertainment_sports", 0.0),
        "business_share_visitor": visitor_shares.get("business_education", 0.0),
        "business_share_tourist": tourist_shares.get("business_education", 0.0),
        "medical_share_visitor": visitor_shares.get("medical_wellness", 0.0),
        "medical_share_tourist": tourist_shares.get("medical_wellness", 0.0),
        "religious_share_visitor": visitor_shares.get("religious", 0.0),
        "religious_share_tourist": tourist_shares.get("religious", 0.0),
        "incentive_others_share_visitor": visitor_shares.get("incentive_others", 0.0),
        "incentive_others_share_tourist": tourist_shares.get("incentive_others", 0.0),
    }
    return record


def parse_national_top_destinations_and_districts(year: int) -> Tuple[List[Dict], List[Dict]]:
    """Parse Jadual 8A (Top Destinations) and 8B (Top Districts) from National DTS."""
    nat_files = glob.glob(str(DTS_BASE_DIR / f"{year}/national/*.xlsx"))
    if not nat_files:
        return [], []

    file_path = Path(nat_files[0])
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  [Error] Cannot load national {year}: {e}")
        return [], []

    destinations_list = []
    districts_list = []

    # Sheet 8A (Destinations)
    sheet_8a = next((s for s in wb.sheetnames if "8a" in s.lower() or s.strip() == "8A"), None)
    if sheet_8a:
        ws = wb[sheet_8a]
        rows_8a = list(ws.iter_rows(min_row=1, max_row=35, max_col=10, values_only=True))
        for r_idx in range(3, len(rows_8a)):
            row_vals = rows_8a[r_idx]
            for c_state, c_dest in [(0, 1), (1, 2), (3, 4), (5, 6), (6, 7)]:
                if c_state < len(row_vals) and c_dest < len(row_vals):
                    st_val = row_vals[c_state]
                    dest_val = row_vals[c_dest]
                    if st_val and dest_val and isinstance(st_val, str) and not st_val.strip().startswith("*"):
                        st_clean = st_val.strip()
                        if "negeri" in st_clean.lower() and "state" in st_clean.lower():
                            continue
                        m_key = match_state_from_name(st_clean)
                        if m_key:
                            meta = STATE_METADATA[m_key]
                            dests = [d.strip() for d in str(dest_val).split("\n") if d.strip() and not d.strip().startswith("*")]
                            for rank, d_name in enumerate(dests, 1):
                                destinations_list.append({
                                    "year": year,
                                    "state": meta["name"],
                                    "state_code": meta["code"],
                                    "region": meta["region"],
                                    "rank": rank,
                                    "destination_name": d_name,
                                })

    # Sheet 8B (Districts)
    sheet_8b = next((s for s in wb.sheetnames if "8b" in s.lower() or s.strip() == "8B"), None)
    if sheet_8b:
        ws = wb[sheet_8b]
        rows_8b = list(ws.iter_rows(min_row=1, max_row=35, max_col=10, values_only=True))
        for r_idx in range(3, len(rows_8b)):
            row_vals = rows_8b[r_idx]
            for c_state, c_dist in [(0, 1), (1, 2), (3, 4), (5, 6), (6, 7)]:
                if c_state < len(row_vals) and c_dist < len(row_vals):
                    st_val = row_vals[c_state]
                    dist_val = row_vals[c_dist]
                    if st_val and dist_val and isinstance(st_val, str) and not st_val.strip().startswith("*"):
                        st_clean = st_val.strip()
                        if "negeri" in st_clean.lower() and "state" in st_clean.lower():
                            continue
                        m_key = match_state_from_name(st_clean)
                        if m_key:
                            meta = STATE_METADATA[m_key]
                            districts = [d.strip() for d in str(dist_val).split("\n") if d.strip() and not d.strip().startswith("*")]
                            for rank, d_name in enumerate(districts, 1):
                                districts_list.append({
                                    "year": year,
                                    "state": meta["name"],
                                    "state_code": meta["code"],
                                    "region": meta["region"],
                                    "rank": rank,
                                    "district_name": d_name,
                                })

    wb.close()

    # Deduplicate in case overlapping column probes captured same pair
    def dedupe(lst, key_cols):
        seen = set()
        out = []
        for row in lst:
            key = tuple(row[k] for k in key_cols)
            if key not in seen:
                seen.add(key)
                out.append(row)
        return out

    destinations_clean = dedupe(destinations_list, ["year", "state", "rank", "destination_name"])
    districts_clean = dedupe(districts_list, ["year", "state", "rank", "district_name"])

    return destinations_clean, districts_clean


def run_pipeline():
    print("=" * 70)
    print("Parsing DTS Purpose of Visit & Top Destinations/Districts (2018–2025)")
    print("=" * 70)

    years = sorted([int(p.name) for p in DTS_BASE_DIR.glob("*") if p.is_dir() and p.name.isdigit()])
    print(f"Discovered {len(years)} DTS survey years: {years}\n")

    all_purpose_records = []
    all_destinations = []
    all_districts = []

    from concurrent.futures import ProcessPoolExecutor

    for y in years:
        # 1. State Purpose of Visit
        state_dir = DTS_BASE_DIR / f"{y}/state"
        state_files = sorted(list(state_dir.glob("*.xlsx")))
        print(f"Year {y}: processing {len(state_files)} state files with ProcessPool...")

        with ProcessPoolExecutor(max_workers=4) as executor:
            futures = [executor.submit(parse_state_jadual_8, f, y) for f in state_files]
            y_records = 0
            for fut in futures:
                rec = fut.result()
                if rec:
                    all_purpose_records.append(rec)
                    y_records += 1

        print(f"  -> Successfully extracted purpose of visit for {y_records} states.")

        # 2. National Top Destinations (8A) and Districts (8B)
        dests, dists = parse_national_top_destinations_and_districts(y)
        all_destinations.extend(dests)
        all_districts.extend(dists)
        print(f"  -> Extracted {len(dests)} destination entries and {len(dists)} district entries from National 8A/8B.")

    # Convert to DataFrames
    df_purpose = pd.DataFrame(all_purpose_records)
    df_dest = pd.DataFrame(all_destinations)
    df_dist = pd.DataFrame(all_districts)

    print("\nPipeline Extraction Summary:")
    print(f"  1. state_purpose_of_visit_panel: {len(df_purpose)} records ({df_purpose['year'].nunique()} years, {df_purpose['state'].nunique()} states)")
    print(f"  2. state_top_destinations:       {len(df_dest)} records")
    print(f"  3. state_top_districts:          {len(df_dist)} records")

    # Export to CSV
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    df_purpose.to_csv(PROCESSED_DIR / "state_purpose_of_visit_panel.csv", index=False)
    df_dest.to_csv(PROCESSED_DIR / "state_top_destinations.csv", index=False)
    df_dist.to_csv(PROCESSED_DIR / "state_top_districts.csv", index=False)

    # Insert / Replace in DuckDB and write Parquet via DuckDB native engine
    con = duckdb.connect(str(DUCKDB_PATH))
    con.execute("CREATE OR REPLACE TABLE state_purpose_of_visit_panel AS SELECT * FROM df_purpose")
    con.execute("CREATE OR REPLACE TABLE state_top_destinations AS SELECT * FROM df_dest")
    con.execute("CREATE OR REPLACE TABLE state_top_districts AS SELECT * FROM df_dist")

    con.execute(f"COPY state_purpose_of_visit_panel TO '{PROCESSED_DIR}/state_purpose_of_visit_panel.parquet' (FORMAT PARQUET)")
    con.execute(f"COPY state_top_destinations TO '{PROCESSED_DIR}/state_top_destinations.parquet' (FORMAT PARQUET)")
    con.execute(f"COPY state_top_districts TO '{PROCESSED_DIR}/state_top_districts.parquet' (FORMAT PARQUET)")

    print("\nDuckDB Tables & Parquet Files Created Successfully:")
    for tbl in ["state_purpose_of_visit_panel", "state_top_destinations", "state_top_districts"]:
        n = con.execute(f"SELECT COUNT(*) FROM {tbl}").fetchone()[0]
        print(f"  ✓ {tbl:<35} {n:>6d} rows")
    con.close()
    print("=" * 70)
    print("=" * 70)


if __name__ == "__main__":
    run_pipeline()
