"""
TSA Data Ingestion and Parser
Extracts tidy time-series data from data/tsa/tourism_2025.xlsx (2015-2025).
Standardizes product categories, separates estimate/preliminary flags,
and outputs tidy tables according to AGENTS.md specifications.
"""

import re
from pathlib import Path
from typing import Dict, List, Tuple
import openpyxl
import pandas as pd
from src.config.paths import TSA_DIR
from src.ingestion.numeric import source_number
from src.analytics.accounting import calc_vai

TSA_FILE_PATH = TSA_DIR / "tourism_2025.xlsx"

# Canonical English naming for the 8 TSA characteristic product categories
CANONICAL_PRODUCTS = [
    {
        "id": "accommodation",
        "product": "Accommodation services",
        "industry": "Accommodation services",
        "row_jad4": 6,
        "row_jad5": 6,
        "row_jad6_supply": 6,
        "row_jad6_ratio": 17,
        "row_jad7": 6,
    },
    {
        "id": "food_beverage",
        "product": "Food and beverage serving services",
        "industry": "Food and beverage serving services",
        "row_jad4": 7,
        "row_jad5": 7,
        "row_jad6_supply": 7,
        "row_jad6_ratio": 18,
        "row_jad7": 7,
    },
    {
        "id": "passenger_transport",
        "product": "Passenger transport services",
        "industry": "Passenger transport services",
        "row_jad4": 8,
        "row_jad5": 8,
        "row_jad6_supply": 8,
        "row_jad6_ratio": 19,
        "row_jad7": 8,
    },
    {
        "id": "travel_agency",
        "product": "Travel agencies and other reservation services",
        "industry": "Travel agencies and other reservation services",
        "row_jad4": 9,
        "row_jad5": 9,
        "row_jad6_supply": 9,
        "row_jad6_ratio": 20,
        "row_jad7": 9,
    },
    {
        "id": "cultural_sports_recreation",
        "product": "Cultural, sports and recreational services",
        "industry": "Cultural, sports and recreational services",
        "row_jad4": 10,
        "row_jad5": 10,
        "row_jad6_supply": 10,
        "row_jad6_ratio": 21,
        "row_jad7": 10,
    },
    {
        "id": "automotive_fuel",
        "product": "Retail sale of automotive fuel",
        "industry": "Retail sale of automotive fuel",
        "row_jad4": 11,
        "row_jad5": 11,
        "row_jad6_supply": 11,
        "row_jad6_ratio": 22,
        "row_jad7": 11,
    },
    {
        "id": "country_specific_goods",
        "product": "Country-specific tourism characteristic goods",
        "industry": "Retail trade of tourism characteristic goods",
        "row_jad4": 12,
        "row_jad5": 12,
        "row_jad6_supply": 12,
        "row_jad6_ratio": 23,
        "row_jad7": 12,
    },
    {
        "id": "country_specific_services",
        "product": "Country-specific tourism characteristic services",
        "industry": "Country-specific tourism characteristic services",
        "row_jad4": 13,
        "row_jad5": 13,
        "row_jad6_supply": 13,
        "row_jad6_ratio": 24,
        "row_jad7": 13,
    },
]


def parse_year_header(raw_header: str) -> Tuple[int, str]:
    """
    Parses headers like '2015', '2024e', '2025p' into (year: int, status: str).
    """
    raw_str = str(raw_header).strip()
    match = re.match(r"^(\d{4})([a-zA-Z]*)$", raw_str)
    if match:
        year = int(match.group(1))
        flag = match.group(2).lower()
        status_map = {"": "actual", "e": "estimate", "p": "preliminary", "r": "revised"}
        return year, status_map.get(flag, flag)
    raise ValueError(f"Unable to parse year from header: {raw_header}")


def ingest_tsa_tables(file_path: Path = TSA_FILE_PATH) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Ingests TSA tables from tourism_2025.xlsx and produces:
    1. tourism_product_year (tidy dataframe of 8 products x 11 years)
    2. tsa_macro_year (macro TDGVA, TDGDP, total employment, etc. x 11 years)
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws_jad4 = wb["Jad 4"]  # ITC
    ws_jad5 = wb["Jad 5"]  # GVA
    ws_jad6 = wb["Jad 6"]  # Supply, Tourism Ratio, TDGVA
    ws_jad7 = wb["Jad 7"]  # Employment

    # Discover column indices and years (row 3 has year headers in col 2 to 12)
    col_mapping = []  # list of (col_idx, year, status)
    for col in range(2, 13):
        val = ws_jad6.cell(3, col).value
        if val is not None:
            year, status = parse_year_header(str(val))
            col_mapping.append((col, year, status))

    records = []
    for prod in CANONICAL_PRODUCTS:
        for col_idx, year, status in col_mapping:
            # 1. Domestic Supply (Jad 6, rows 6-13)
            supply_val = source_number(ws_jad6.cell(prod["row_jad6_supply"], col_idx).value)

            # 2. Gross Value Added (Jad 5, rows 6-13)
            gva_val = source_number(ws_jad5.cell(prod["row_jad5"], col_idx).value)

            # 3. Internal Tourism Consumption (Jad 4, rows 6-13)
            itc_val = source_number(ws_jad4.cell(prod["row_jad4"], col_idx).value)

            # 4. Tourism Ratio (Jad 6, rows 17-24)
            ratio_val = source_number(ws_jad6.cell(prod["row_jad6_ratio"], col_idx).value)

            # 5. Employment in Thousand Persons (Jad 7, rows 6-13)
            emp_val = source_number(ws_jad7.cell(prod["row_jad7"], col_idx).value)

            # Economic computations per AGENTS.md
            # Value-Added Intensity (VAI) = GVA / DomesticSupply
            vai = calc_vai(gva_val, supply_val)

            # Analytical proxy: Estimated tourism-attributable GVA = ITC * VAI
            # (or equivalently: GVA * TourismRatio)
            estimated_tourism_gva = itc_val * vai

            # Period classification
            if year <= 2019:
                period = "Pre-COVID (2015-2019)"
            elif year <= 2022:
                period = "Disruption & Recovery (2020-2022)"
            else:
                period = "Post-Recovery (2023-2025)"

            records.append({
                "year": year,
                "period": period,
                "product_id": prod["id"],
                "product": prod["product"],
                "industry": prod["industry"],
                "domestic_supply": round(supply_val, 2),
                "gva": round(gva_val, 2),
                "itc": round(itc_val, 2),
                "tourism_ratio": round(ratio_val, 4),
                "employment_thousands": round(emp_val, 2),
                "vai": round(vai, 4),
                "estimated_tourism_gva": round(estimated_tourism_gva, 2),
                "data_status": status,
            })

    df_product = pd.DataFrame(records)

    # Ingest Macro TSA aggregates (Rows in Jad 6 and Jad 5)
    macro_records = []
    for col_idx, year, status in col_mapping:
        tdgva = source_number(ws_jad6.cell(28, col_idx).value)
        tdgdp = source_number(ws_jad6.cell(29, col_idx).value)
        tdgva_share_gva = source_number(ws_jad6.cell(36, col_idx).value)
        tdgdp_share_gdp = source_number(ws_jad6.cell(37, col_idx).value)

        # Macro totals
        total_supply = source_number(ws_jad6.cell(14, col_idx).value)
        total_gvati = source_number(ws_jad5.cell(14, col_idx).value)
        total_itc = source_number(ws_jad4.cell(14, col_idx).value)
        total_emp = source_number(ws_jad7.cell(14, col_idx).value)
        overall_ratio = source_number(ws_jad6.cell(25, col_idx).value)

        macro_records.append({
            "year": year,
            "data_status": status,
            "tdgva": round(tdgva, 2),
            "tdgdp": round(tdgdp, 2),
            "tdgva_share_gva": round(tdgva_share_gva, 2),
            "tdgdp_share_gdp": round(tdgdp_share_gdp, 2),
            "total_domestic_supply": round(total_supply, 2),
            "total_gvati": round(total_gvati, 2),
            "total_itc": round(total_itc, 2),
            "total_employment_thousands": round(total_emp, 2),
            "overall_tourism_ratio": round(overall_ratio, 4),
        })

    df_macro = pd.DataFrame(macro_records)

    return df_product, df_macro


if __name__ == "__main__":
    df_prod, df_macro = ingest_tsa_tables()
    print("=== Tourism Product Year (Head) ===")
    print(df_prod.head(10)[["year", "product", "domestic_supply", "gva", "itc", "vai", "estimated_tourism_gva"]])
    print("\n=== TSA Macro Year ===")
    print(df_macro[["year", "tdgva", "tdgdp", "tdgva_share_gva", "total_itc"]])
