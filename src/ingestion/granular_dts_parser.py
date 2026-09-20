"""
Granular DTS Sub-Table Parser (2025).
Extracts structural diagnostic and exploratory drivers across all 16 Malaysian States & FTs:
  - Jadual 12: Types of Accommodation (Paid Commercial vs. Unpaid VFR)
  - Jadual 8:  Purpose of Visit (Holiday/Leisure vs. VFR vs. Business)
  - Jadual 11: Mode of Transport (Private Vehicle vs. Air vs. Bus/Train)
  - Jadual 13a: Household Income Demographics (B40, M40, T20, Affluence Index)
  - Jadual 14: Hotel & Room Inventory (Supply Capacity & Star Ratings)

Missing Value Policy (IMPLEMENTATION_PLAN.md Sprint 1 & AGENTS.md Rule 6):
Missing observations propagate as np.nan / None (SQL NULL), never coerced to 0.0.

Exports 'state_granular_profile' to DuckDB and Parquet.
"""

import re
import sys
from pathlib import Path
from typing import Dict, List, Optional
import duckdb
import numpy as np
import openpyxl
import pandas as pd

from src.config.paths import PROJECT_ROOT, PROCESSED_DATA_DIR, DUCKDB_PATH, DATA_DIR

STATE_DIR_2025 = DATA_DIR / "dts/2025/state"

STATE_METADATA: Dict[str, Dict] = {
    "JOHOR": {"name": "Johor", "code": "MY-01", "region": "Southern"},
    "KEDAH": {"name": "Kedah", "code": "MY-02", "region": "Northern"},
    "KELANTAN": {"name": "Kelantan", "code": "MY-03", "region": "East Coast"},
    "MELAKA": {"name": "Melaka", "code": "MY-04", "region": "Southern"},
    "NEGERI SEMBILAN": {"name": "Negeri Sembilan", "code": "MY-05", "region": "Central"},
    "PAHANG": {"name": "Pahang", "code": "MY-06", "region": "East Coast"},
    "PULAU PINANG": {"name": "Pulau Pinang", "code": "MY-07", "region": "Northern"},
    "PERAK": {"name": "Perak", "code": "MY-08", "region": "Northern"},
    "PERLIS": {"name": "Perlis", "code": "MY-09", "region": "Northern"},
    "SELANGOR": {"name": "Selangor", "code": "MY-10", "region": "Central"},
    "TERENGGANU": {"name": "Terengganu", "code": "MY-11", "region": "East Coast"},
    "SABAH": {"name": "Sabah", "code": "MY-12", "region": "East Malaysia"},
    "SARAWAK": {"name": "Sarawak", "code": "MY-13", "region": "East Malaysia"},
    "W.P. KUALA LUMPUR": {"name": "W.P. Kuala Lumpur", "code": "MY-14", "region": "Central"},
    "W.P. LABUAN": {"name": "W.P. Labuan", "code": "MY-15", "region": "East Malaysia"},
    "W.P. PUTRAJAYA": {"name": "W.P. Putrajaya", "code": "MY-16", "region": "Central"},
}


def match_state_from_path(filepath: Path) -> str:
    fname = filepath.name.upper()
    if "KUALA LUMPUR" in fname:
        return "W.P. KUALA LUMPUR"
    if "LABUAN" in fname:
        return "W.P. LABUAN"
    if "PUTRAJAYA" in fname:
        return "W.P. PUTRAJAYA"
    if "PINANG" in fname or "PENANG" in fname:
        return "PULAU PINANG"
    if "SEMBILAN" in fname:
        return "NEGERI SEMBILAN"
    for k in STATE_METADATA:
        if k in fname:
            return k
    raise ValueError(f"Could not identify state from filename: {filepath.name}")


def clean_num(val) -> Optional[float]:
    """
    Cleans cell value to float. Returns None for empty, dashes, or unparseable text.
    Preserves actual zeros (e.g. 0 or 0.0).
    """
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val) if not np.isnan(val) else None
    val_str = str(val).replace(",", "").strip()
    if not val_str or val_str == "-" or val_str.lower() in ["na", "n/a", "null", "nil"]:
        return None
    match = re.search(r"[-+]?\d*\.\d+|\d+", val_str)
    return float(match.group()) if match else None


def safe_round(val: Optional[float], decimals: int = 2) -> Optional[float]:
    """Rounds float if present and valid; returns None for NaN/None."""
    if val is None or pd.isna(val) or np.isnan(val):
        return None
    return round(float(val), decimals)


def parse_state_granular(filepath: Path) -> Dict:
    state_key = match_state_from_path(filepath)
    meta = STATE_METADATA[state_key]

    wb = openpyxl.load_workbook(filepath, data_only=True)

    # 1. Parse Jadual 11 & 12 (Transport Modes & Accommodation Types)
    ws11_12 = wb["Jadual 11 & 12"]

    # --- Mode of Transport (Jadual 11, top half, col 7 is 2025 Tourists %) ---
    air_share = np.nan
    private_veh_share = np.nan
    bus_share = np.nan
    train_share = np.nan

    for r in range(7, 18):
        lbl = str(ws11_12.cell(r, 1).value or "").lower()
        val = clean_num(ws11_12.cell(r, 7).value)
        if val is not None:
            if "udara" in lbl or "air" in lbl:
                air_share = val
            elif "persendirian" in lbl or "private" in lbl:
                private_veh_share = val
            elif "bas" in lbl or "bus" in lbl:
                bus_share = val
            elif "kereta api" in lbl or "train" in lbl:
                train_share = val

    # --- Type of Accommodation (Jadual 12, bottom half, col 6 is 2025 %) ---
    hotel_share = np.nan
    chalet_share = np.nan
    apt_share = np.nan
    homestay_share = np.nan
    rest_house_share = np.nan
    vfr_accom_share = np.nan

    for r in range(20, min(40, ws11_12.max_row + 1)):
        lbl = str(ws11_12.cell(r, 1).value or "").lower()
        val = clean_num(ws11_12.cell(r, 6).value)
        if val is not None:
            if "saudara" in lbl or "relatives" in lbl:
                vfr_accom_share = val
            elif "hotel" in lbl:
                hotel_share = val
            elif "chalet" in lbl:
                chalet_share = val
            elif "apartmen" in lbl or "apartment" in lbl:
                apt_share = val
            elif "inap desa" in lbl or "homestay" in lbl:
                homestay_share = val
            elif "rumah rehat" in lbl or "rest house" in lbl:
                rest_house_share = val

    paid_components = [hotel_share, chalet_share, apt_share, homestay_share, rest_house_share]
    valid_paid = [x for x in paid_components if not pd.isna(x)]
    paid_commercial_share = sum(valid_paid) if valid_paid else np.nan

    # 2. Parse Jadual 8 (Purpose of Visit)
    ws8 = wb["Jadual 8"]
    holiday_share = np.nan
    vfr_purpose_share = np.nan
    shopping_purpose_share = np.nan
    medical_share = np.nan
    business_share = np.nan

    # In Jadual 8: Col 5 has purpose label for tourists, Col 6 has Percentage share (%)
    for r in range(6, min(17, ws8.max_row + 1)):
        lbl = str(ws8.cell(r, 5).value or "").lower()
        val = clean_num(ws8.cell(r, 6).value)
        if val is not None:
            if "percutian" in lbl or "holiday" in lbl or "leisure" in lbl:
                holiday_share = val
            elif "saudara" in lbl or "relatives" in lbl:
                vfr_purpose_share = val
            elif "membeli-belah" in lbl or "shopping" in lbl:
                shopping_purpose_share = val
            elif "perubatan" in lbl or "medical" in lbl or "wellness" in lbl:
                medical_share = val
            elif "rasmi" in lbl or "perniagaan" in lbl or "business" in lbl or "pendidikan" in lbl:
                business_share = val

    # 3. Parse Jadual 13a (Household Income Demographics)
    b40_share = np.nan
    m40_share = np.nan
    t20_share = np.nan

    if "Jadual 13a" in wb.sheetnames:
        ws13a = wb["Jadual 13a"]
        low_inc = []
        for r in range(7, min(15, ws13a.max_row + 1)):
            lbl = str(ws13a.cell(r, 2).value or "").lower()
            val = clean_num(ws13a.cell(r, 4).value)
            if val is not None:
                if "1,000" in lbl or "3,000" in lbl or "5,000" in lbl:
                    low_inc.append(val)
                elif "10,000" in lbl and "≥" not in lbl and ">" not in lbl:
                    m40_share = val
                elif "≥ 10,001" in lbl or "> 10,000" in lbl or "10,001" in lbl:
                    t20_share = val
        if low_inc:
            b40_share = sum(low_inc)

    if not pd.isna(m40_share) and not pd.isna(t20_share):
        affluence_index = round((m40_share * 1.0) + (t20_share * 2.0), 2)
    elif not pd.isna(m40_share):
        affluence_index = round(m40_share * 1.0, 2)
    elif not pd.isna(t20_share):
        affluence_index = round(t20_share * 2.0, 2)
    else:
        affluence_index = None

    # 4. Parse Jadual 14 & 15 (Hotel & Room Inventory)
    total_hotels = None
    total_rooms = None
    star_5_rooms = None
    star_4_rooms = None
    star_3_rooms = None

    if "Jadual 14 & 15" in wb.sheetnames:
        ws14 = wb["Jadual 14 & 15"]
        for r in range(5, min(20, ws14.max_row + 1)):
            lbl = str(ws14.cell(r, 1).value or "").lower()
            h_count = clean_num(ws14.cell(r, 3).value)
            r_count = clean_num(ws14.cell(r, 6).value)

            if "5-bintang" in lbl or "5-star" in lbl:
                star_5_rooms = int(r_count) if r_count is not None else None
            elif "4-bintang" in lbl or "4-star" in lbl:
                star_4_rooms = int(r_count) if r_count is not None else None
            elif "3-bintang" in lbl or "3-star" in lbl:
                star_3_rooms = int(r_count) if r_count is not None else None
            elif "jumlah" in lbl or "total" in lbl:
                total_hotels = int(h_count) if h_count is not None else None
                total_rooms = int(r_count) if r_count is not None else None

    wb.close()

    star_rated_components = [star_5_rooms, star_4_rooms, star_3_rooms]
    valid_star = [x for x in star_rated_components if x is not None]
    star_rated_rooms = sum(valid_star) if valid_star else None

    if total_rooms is not None and total_rooms > 0 and star_rated_rooms is not None:
        star_room_share_pct = round((star_rated_rooms / total_rooms * 100.0), 2)
    else:
        star_room_share_pct = None

    return {
        "state": meta["name"],
        "state_code": meta["code"],
        "region": meta["region"],
        # Accommodation Structure
        "paid_commercial_share_pct": safe_round(paid_commercial_share),
        "unpaid_vfr_share_pct": safe_round(vfr_accom_share),
        "hotel_share_pct": safe_round(hotel_share),
        "homestay_share_pct": safe_round(homestay_share),
        "apartment_share_pct": safe_round(apt_share),
        # Purpose of Visit
        "holiday_leisure_share_pct": safe_round(holiday_share),
        "vfr_purpose_share_pct": safe_round(vfr_purpose_share),
        "shopping_purpose_share_pct": safe_round(shopping_purpose_share),
        "medical_wellness_share_pct": safe_round(medical_share),
        "business_mice_share_pct": safe_round(business_share),
        # Mode of Transport
        "private_vehicle_share_pct": safe_round(private_veh_share),
        "air_transport_share_pct": safe_round(air_share),
        "bus_transport_share_pct": safe_round(bus_share),
        "train_transport_share_pct": safe_round(train_share),
        # Income & Affluence
        "b40_share_pct": safe_round(b40_share),
        "m40_share_pct": safe_round(m40_share),
        "t20_share_pct": safe_round(t20_share),
        "affluence_index": affluence_index,
        # Hotel Supply Capacity
        "total_hotels": total_hotels,
        "total_rooms": total_rooms,
        "star_rated_rooms": star_rated_rooms,
        "star_room_share_pct": star_room_share_pct,
    }


def run_granular_ingestion() -> pd.DataFrame:
    print("Ingesting Granular DTS Sub-Tables for all 16 Malaysian States...")
    files = sorted(list(STATE_DIR_2025.glob("*.xlsx")))
    if not files:
        raise FileNotFoundError(f"No 2025 state files found in {STATE_DIR_2025}")

    records = []
    for f in files:
        rec = parse_state_granular(f)
        records.append(rec)

    df = pd.DataFrame(records)

    con = duckdb.connect(str(DUCKDB_PATH))
    con.execute("CREATE OR REPLACE TABLE state_granular_profile AS SELECT * FROM df")
    con.close()

    parquet_path = PROCESSED_DATA_DIR / "state_granular_profile.parquet"
    df.to_parquet(parquet_path, index=False)

    print(f"Successfully processed {len(df)} states into 'state_granular_profile'.")
    print(f"DuckDB table updated & Parquet saved to {parquet_path}")
    return df


if __name__ == "__main__":
    df = run_granular_ingestion()
    print("\nSample Preview:")
    cols_to_preview = [
        "state",
        "paid_commercial_share_pct",
        "unpaid_vfr_share_pct",
        "affluence_index",
        "star_room_share_pct",
    ]
    print(df[cols_to_preview].head(8).to_string(index=False))
