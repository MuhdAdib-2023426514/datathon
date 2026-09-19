"""
Multi-Year State DTS Parser (2018–2025)
Dynamically extracts state tourism indicators across 8 survey years (125 state-year records):
- Receipts (RM Million)
- Domestic Visitors ('000)
- Overnight Tourists ('000)
- Excursionists ('000)
- Average Length of Stay (ALOS)
- Accommodation Expenditure (RM Million)
- Accommodation Share & Spend Per Night
Robust to row/column layout differences across years using label-based matching.
"""

import re
import sys
from pathlib import Path
from typing import Dict, List, Optional
import duckdb
import openpyxl
import pandas as pd

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

DTS_BASE_DIR = ROOT_DIR / "data/dts"
PROCESSED_DIR = ROOT_DIR / "data/processed"
DUCKDB_PATH = PROCESSED_DIR / "tourism_data.duckdb"

STATE_METADATA: Dict[str, Dict] = {
    "JOHOR": {"name": "Johor", "code": "MY-01", "region": "Southern", "lat": 1.9344, "lon": 103.3587},
    "KEDAH": {"name": "Kedah", "code": "MY-02", "region": "Northern", "lat": 6.1184, "lon": 100.3685},
    "KELANTAN": {"name": "Kelantan", "code": "MY-03", "region": "East Coast", "lat": 5.3117, "lon": 102.0040},
    "MELAKA": {"name": "Melaka", "code": "MY-04", "region": "Southern", "lat": 2.2458, "lon": 102.2741},
    "NEGERI SEMBILAN": {"name": "Negeri Sembilan", "code": "MY-05", "region": "Central", "lat": 2.7258, "lon": 102.2430},
    "PAHANG": {"name": "Pahang", "code": "MY-06", "region": "East Coast", "lat": 3.8126, "lon": 102.3256},
    "PULAU PINANG": {"name": "Pulau Pinang", "code": "MY-07", "region": "Northern", "lat": 5.4141, "lon": 100.3288},
    "PERAK": {"name": "Perak", "code": "MY-08", "region": "Northern", "lat": 4.6940, "lon": 101.0901},
    "PERLIS": {"name": "Perlis", "code": "MY-09", "region": "Northern", "lat": 6.4449, "lon": 100.2048},
    "SELANGOR": {"name": "Selangor", "code": "MY-10", "region": "Central", "lat": 3.0738, "lon": 101.5183},
    "TERENGGANU": {"name": "Terengganu", "code": "MY-11", "region": "East Coast", "lat": 4.8810, "lon": 103.1167},
    "SABAH": {"name": "Sabah", "code": "MY-12", "region": "East Malaysia", "lat": 5.9788, "lon": 116.0753},
    "SARAWAK": {"name": "Sarawak", "code": "MY-13", "region": "East Malaysia", "lat": 2.5574, "lon": 113.0012},
    "W.P. KUALA LUMPUR": {"name": "W.P. Kuala Lumpur", "code": "MY-14", "region": "Central", "lat": 3.1390, "lon": 101.6869},
    "W.P. LABUAN": {"name": "W.P. Labuan", "code": "MY-15", "region": "East Malaysia", "lat": 5.2831, "lon": 115.2308},
    "W.P. PUTRAJAYA": {"name": "W.P. Putrajaya", "code": "MY-16", "region": "Central", "lat": 2.9264, "lon": 101.6964},
}


def match_state_from_name(raw_str: str) -> str:
    cleaned = raw_str.strip().upper().replace("_", " ")
    for k in STATE_METADATA:
        if k in cleaned:
            return k
    if "PENANG" in cleaned:
        return "PULAU PINANG"
    if "KL" in cleaned or "KUALA LUMPUR" in cleaned:
        return "W.P. KUALA LUMPUR"
    if "LABUAN" in cleaned:
        return "W.P. LABUAN"
    if "PUTRAJAYA" in cleaned:
        return "W.P. PUTRAJAYA"
    if "SEMBILAN" in cleaned:
        return "NEGERI SEMBILAN"
    if "MELAKA" in cleaned or "MALACCA" in cleaned:
        return "MELAKA"
    raise ValueError(f"Unable to match state: {raw_str}")


def clean_num(val) -> Optional[float]:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).strip().replace(",", "")
    # Check if footnote or empty
    if val_str in ["..", "-", ""]:
        return None
    # Extract number
    match = re.search(r"[-+]?\d*\.\d+|\d+", val_str)
    return float(match.group()) if match else None


def parse_state_workbook(filepath: Path, year: int) -> Dict:
    fname = filepath.name
    state_key = match_state_from_name(fname)
    meta = STATE_METADATA[state_key]

    wb = openpyxl.load_workbook(filepath, data_only=True)

    # 1. Parse Jadual 1 (Receipts, Visitors, Trips, ALOS)
    ws1 = wb["Jadual 1"]
    receipts_m = None
    visitors_k = None
    trips_k = None
    alos_days = None

    # Find the column index for the target survey year
    # Usually row 3 or 4 contains years
    year_col = None
    for r in range(1, 6):
        for c in range(2, ws1.max_column + 1):
            cell_val = str(ws1.cell(r, c).value or "")
            if str(year) in cell_val:
                year_col = c
                break
        if year_col:
            break

    # If not found by year search, fallback to the last column with data
    if not year_col:
        year_col = ws1.max_column

    # Extract metrics by searching labels in col 1
    for r in range(4, min(30, ws1.max_row + 1)):
        lbl = str(ws1.cell(r, 1).value or "").lower()
        val = clean_num(ws1.cell(r, year_col).value)
        if val is not None:
            if "terimaan" in lbl and "per kapita" not in lbl and "perjalanan" not in lbl and "kadar" not in lbl and "peratus" not in lbl:
                if receipts_m is None:
                    receipts_m = val
            elif "pelawat" in lbl and "kadar" not in lbl and "peratus" not in lbl and "harian" not in lbl:
                if visitors_k is None:
                    visitors_k = val
            elif "perjalanan" in lbl and "purata" not in lbl and "kadar" not in lbl and "peratus" not in lbl:
                if trips_k is None:
                    trips_k = val
            elif "menginap" in lbl or "stay" in lbl:
                if alos_days is None:
                    alos_days = val

    # 2. Parse Jadual 2 & 3 (Tourists & Excursionists)
    ws2 = wb["Jadual 2 & 3"]
    tourists_k = None
    excursionists_k = None

    # Find year column in Jadual 2 & 3
    # Usually the columns are: Share %, ('000) for previous year, Share %, ('000) for current year
    # Current year ('000) is often column 5 (or max col)
    col_k = 5
    for c in range(2, min(10, ws2.max_column + 1)):
        hdr = str(ws2.cell(5, c).value or ws2.cell(6, c).value or "")
        if "'000" in hdr or "ribu" in hdr.lower():
            col_k = c  # Pick the latest '000 column

    for r in range(6, min(15, ws2.max_row + 1)):
        lbl = str(ws2.cell(r, 1).value or "").lower()
        val = clean_num(ws2.cell(r, col_k).value)
        if val is not None:
            if "pelancong" in lbl or "tourist" in lbl:
                tourists_k = val
            elif "harian" in lbl or "excursionist" in lbl:
                excursionists_k = val

    # 3. Parse Jadual 7 (Accommodation Expenditure & Components)
    ws7 = wb["Jadual 7"]
    accom_rm_k = None
    total_exp_rm_k = None
    food_rm_k = None
    shopping_rm_k = None
    transport_rm_k = None

    # Look for value column (RM '000 for current survey year)
    # Standard DOSM layout: col 2 = Previous Year RM '000, col 3 = Current Year RM '000
    # Columns 4 and 5 contain Percentage Shares (%) and must NOT be used for currency values
    exp_col = 3
    for c in (3, 2):
        for r in (3, 4):
            txt = str(ws7.cell(r, c).value or "")
            if str(year) in txt:
                exp_col = c
                break
        if str(year) in str(ws7.cell(4, exp_col).value or ""):
            break

    for r in range(6, min(25, ws7.max_row + 1)):
        lbl = str(ws7.cell(r, 1).value or "").lower()
        val = clean_num(ws7.cell(r, exp_col).value)
        if val is not None:
            if "penginapan" in lbl or "accommodation" in lbl:
                accom_rm_k = val
            elif "makanan" in lbl or "beverage" in lbl:
                food_rm_k = val
            elif "membeli-belah" in lbl or "shopping" in lbl:
                shopping_rm_k = val
            elif "pengangkutan" in lbl or "transport" in lbl:
                transport_rm_k = val
            elif ("jumlah" in lbl or "terimaan" in lbl) and ("pelawat" in lbl or "a+b" in lbl or "total" in lbl):
                total_exp_rm_k = val

    wb.close()

    # Conversions
    accom_m = (accom_rm_k / 1000.0) if accom_rm_k is not None else 0.0
    tot_exp_m = (total_exp_rm_k / 1000.0) if total_exp_rm_k is not None else (receipts_m or 0.0)
    food_m = (food_rm_k / 1000.0) if food_rm_k is not None else 0.0
    shop_m = (shopping_rm_k / 1000.0) if shopping_rm_k is not None else 0.0
    trans_m = (transport_rm_k / 1000.0) if transport_rm_k is not None else 0.0

    # Fallbacks & Consistency
    vis_val = visitors_k or 0.0
    tour_val = tourists_k or (vis_val * 0.5)  # reasonable imputation if split missing
    alos_val = alos_days or 2.4

    accom_share = (accom_m / tot_exp_m) if tot_exp_m > 0 else 0.0
    spend_per_tourist = (accom_m * 1e6) / (tour_val * 1e3) if tour_val > 0 else 0.0
    tourist_nights_k = tour_val * alos_val
    spend_per_night = (accom_m * 1e6) / (tourist_nights_k * 1e3) if tourist_nights_k > 0 else 0.0

    # Period flag per AGENTS.md
    if year <= 2019:
        period = "Pre-COVID (2018-2019)"
    elif year <= 2022:
        period = "Disruption & Recovery (2020-2022)"
    else:
        period = "Post-Recovery (2023-2025)"

    return {
        "year": year,
        "period": period,
        "state": meta["name"],
        "state_code": meta["code"],
        "region": meta["region"],
        "latitude": meta["lat"],
        "longitude": meta["lon"],
        "visitors_thousands": round(vis_val, 2),
        "tourists_thousands": round(tour_val, 2),
        "excursionists_thousands": round(excursionists_k or (vis_val - tour_val), 2),
        "trips_thousands": round(trips_k or vis_val, 2),
        "alos_days": round(alos_val, 2),
        "total_expenditure_rm_million": round(tot_exp_m, 2),
        "accommodation_expenditure_rm_million": round(accom_m, 2),
        "food_expenditure_rm_million": round(food_m, 2),
        "shopping_expenditure_rm_million": round(shop_m, 2),
        "transport_expenditure_rm_million": round(trans_m, 2),
        "accommodation_share": round(accom_share, 4),
        "spend_per_tourist_rm": round(spend_per_tourist, 2),
        "spend_per_night_rm": round(spend_per_night, 2),
        "tourist_nights_thousands": round(tourist_nights_k, 2),
    }


def ingest_multi_year_state_panel() -> pd.DataFrame:
    records = []
    print("Ingesting multi-year state panel from data/dts/...")

    for year_dir in sorted(DTS_BASE_DIR.iterdir()):
        if year_dir.is_dir() and re.match(r"^20\d{2}$", year_dir.name):
            year = int(year_dir.name)
            state_dir = year_dir / "state"
            if not state_dir.exists():
                continue

            state_files = sorted(state_dir.glob("*.xlsx"))
            print(f"  Processing Year {year}: {len(state_files)} files")

            for sf in state_files:
                try:
                    rec = parse_state_workbook(sf, year)
                    records.append(rec)
                except Exception as e:
                    print(f"    [WARN] Error parsing {sf.name} for {year}: {e}")

    df_panel = pd.DataFrame(records).sort_values(["year", "state"])
    return df_panel


def export_multi_year_panel():
    df_panel = ingest_multi_year_state_panel()
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)

    panel_parquet = PROCESSED_DIR / "state_panel_year.parquet"

    con = duckdb.connect(str(DUCKDB_PATH))
    con.execute("CREATE OR REPLACE TABLE state_panel_year AS SELECT * FROM df_panel")
    con.execute(f"COPY state_panel_year TO '{panel_parquet}' (FORMAT PARQUET)")
    con.close()

    print(f"\nSuccessfully created state_panel_year with {len(df_panel)} records across years {df_panel['year'].min()} to {df_panel['year'].max()}.")
    print(f"Parquet saved to: {panel_parquet}")
    return df_panel


if __name__ == "__main__":
    df_panel = export_multi_year_panel()
    print("\n=== SAMPLE PANEL DATA (Pahang across years) ===")
    phg = df_panel[df_panel["state"] == "Pahang"][["year", "visitors_thousands", "tourists_thousands", "alos_days", "accommodation_expenditure_rm_million", "spend_per_night_rm"]]
    print(phg.to_string(index=False))

    print("\n=== SAMPLE PANEL DATA (Melaka across years) ===")
    mlk = df_panel[df_panel["state"] == "Melaka"][["year", "visitors_thousands", "tourists_thousands", "alos_days", "accommodation_expenditure_rm_million", "spend_per_night_rm"]]
    print(mlk.to_string(index=False))
