"""
DTS Tourist Income Profile (Jadual 13a) and Hotel Star Inventory (Jadual 14/15) Parser (2018–2025)
Extracts:
1. Tourist monthly household income distribution (B40, M40, T20 shares and Affluence Index)
2. Hotel and room supply by star rating (5-star, 4-star, 3-star, 2-star, 1-star, Orchid, Unrated)
3. Computes luxury room share (% 4-5 star) and budget/unrated share

Exports to DuckDB:
- state_tourist_income_panel
- state_hotel_star_inventory
"""

import glob
import os
import re
import sys
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import duckdb
import openpyxl
import pandas as pd

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

DTS_BASE_DIR = ROOT_DIR / "data/dts"
PROCESSED_DIR = ROOT_DIR / "data/processed"
DUCKDB_PATH = PROCESSED_DIR / "tourism_data.duckdb"

STATE_METADATA: Dict[str, Dict] = {
    "JOHOR": {"name": "Johor", "code": "MY-01", "region": "Southern"},
    "KEDAH": {"name": "Kedah", "code": "MY-02", "region": "Northern"},
    "KELANTAN": {"name": "Kelantan", "code": "MY-03", "region": "East Coast"},
    "MELAKA": {"name": "Melaka", "code": "MY-04", "region": "Southern"},
    "NEGERI SEMBILAN": {"name": "Negeri Sembilan", "code": "MY-05", "region": "Central"},
    "PAHANG": {"name": "Pahang", "code": "MY-06", "region": "East Coast"},
    "PULAU PINANG": {"name": "Pulau Pinang", "code": "MY-07", "region": "Northern"},
    "PERAK": {"name": "Perak", "code": "MY-08", "region": "Northern"},
    "PERLIS": {"name": "Perlis", "code": "MY-09", "region": "Northern"},
    "SELANGOR": {"name": "Selangor", "code": "MY-10", "region": "Central"},
    "TERENGGANU": {"name": "Terengganu", "code": "MY-11", "region": "East Coast"},
    "SABAH": {"name": "Sabah", "code": "MY-12", "region": "East Malaysia"},
    "SARAWAK": {"name": "Sarawak", "code": "MY-13", "region": "East Malaysia"},
    "W.P. KUALA LUMPUR": {"name": "W.P. Kuala Lumpur", "code": "MY-14", "region": "Central"},
    "W.P. LABUAN": {"name": "W.P. Labuan", "code": "MY-15", "region": "East Malaysia"},
    "W.P. PUTRAJAYA": {"name": "W.P. Putrajaya", "code": "MY-16", "region": "Central"},
}


def match_state_from_name(raw_str: str) -> Optional[str]:
    cleaned = raw_str.strip().upper().replace("_", " ").replace(".", "")
    if "PULAU PINANG" in cleaned or "PENANG" in cleaned or "P PINANG" in cleaned:
        return "PULAU PINANG"
    if "KUALA LUMPUR" in cleaned or "KL" in cleaned:
        return "W.P. KUALA LUMPUR"
    if "PUTRAJAYA" in cleaned:
        return "W.P. PUTRAJAYA"
    if "LABUAN" in cleaned:
        return "W.P. LABUAN"
    if "NEGERI SEMBILAN" in cleaned or "N SEMBILAN" in cleaned:
        return "NEGERI SEMBILAN"
    for k in STATE_METADATA:
        if k.replace(".", "") in cleaned:
            return k
    return None


def parse_single_workbook(file_path: Path, year: int) -> Tuple[Optional[Dict], Optional[Dict]]:
    """Parse Jadual 13a (Income) and Jadual 14 (Hotels) for a single state workbook."""
    file_stem = file_path.stem.replace(f"TABLE OF PUBLICATION DTS {year} ", "")
    matched_key = match_state_from_name(file_stem)
    if not matched_key:
        return None, None

    meta = STATE_METADATA[matched_key]

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        return None, None

    # --- 1. JADUAL 13a (TOURIST INCOME DISTRIBUTION) ---
    inc_rec = None
    sheet_13a = next((s for s in wb.sheetnames if "13a" in s.lower() or "13(2)" in s.lower()), None)
    if sheet_13a:
        ws13 = wb[sheet_13a]
        rows13 = list(ws13.iter_rows(min_row=1, max_row=16, max_col=10, values_only=True))

        # Detect target year column index
        yr_col = None
        for r_idx in range(2, min(6, len(rows13))):
            for c_idx, val in enumerate(rows13[r_idx]):
                if val == year or str(val).strip() == str(year):
                    yr_col = c_idx
                    break
            if yr_col is not None:
                break
        if yr_col is None:
            yr_col = 3  # standard fallback

        vals = {}
        for r in rows13:
            label = str(r[1] or r[0] or "")
            raw_v = r[yr_col] if yr_col < len(r) else None
            try:
                fv = float(raw_v) if raw_v is not None else None
            except:
                fv = None

            if fv is not None:
                if "≤ 1,000" in label or "<= 1000" in label or "1,000 kebawah" in label.lower():
                    vals["under_1k"] = fv
                elif "1,001 - 3,000" in label:
                    vals["1k_3k"] = fv
                elif "3,001 - 5,000" in label:
                    vals["3k_5k"] = fv
                elif "5,001 - 10,000" in label:
                    vals["5k_10k"] = fv
                elif "≥ 10,001" in label or ">= 10001" in label or "10,001 ke atas" in label.lower():
                    vals["above_10k"] = fv

        u1k = vals.get("under_1k", 0.0)
        i1_3k = vals.get("1k_3k", 0.0)
        i3_5k = vals.get("3k_5k", 0.0)
        i5_10k = vals.get("5k_10k", 0.0)
        a10k = vals.get("above_10k", 0.0)

        b40 = round(u1k + i1_3k + i3_5k, 2)
        m40 = round(i5_10k, 2)
        t20 = round(a10k, 2)
        affluence = round(m40 + t20, 2)

        inc_rec = {
            "year": year,
            "state": meta["name"],
            "state_code": meta["code"],
            "region": meta["region"],
            "income_under_1k_pct": u1k,
            "income_1k_3k_pct": i1_3k,
            "income_3k_5k_pct": i3_5k,
            "income_5k_10k_pct": i5_10k,
            "income_above_10k_pct": a10k,
            "b40_share_pct": b40,
            "m40_share_pct": m40,
            "t20_share_pct": t20,
            "affluence_index": affluence,
        }

    # --- 2. JADUAL 14 (HOTEL STAR RATING INVENTORY) ---
    hot_rec = None
    sheet_14 = next((s for s in wb.sheetnames if "14" in s), None)
    if sheet_14:
        ws14 = wb[sheet_14]
        rows14 = list(ws14.iter_rows(min_row=1, max_row=22, max_col=10, values_only=True))

        star_data = {
            "h5": 0, "r5": 0,
            "h4": 0, "r4": 0,
            "h3": 0, "r3": 0,
            "h2": 0, "r2": 0,
            "h1": 0, "r1": 0,
            "horchid": 0, "rorchid": 0,
            "hunrated": 0, "runrated": 0,
            "htotal": 0, "rtotal": 0,
        }

        for r in rows14:
            label = str(r[0] or r[1] or "").lower()
            nums = [v for v in r if isinstance(v, (int, float))]
            if len(nums) >= 2:
                h_cnt = int(nums[0])
                r_cnt = int(nums[-1])
                if "5-bintang" in label or "5-star" in label:
                    star_data["h5"], star_data["r5"] = h_cnt, r_cnt
                elif "4-bintang" in label or "4-star" in label:
                    star_data["h4"], star_data["r4"] = h_cnt, r_cnt
                elif "3-bintang" in label or "3-star" in label:
                    star_data["h3"], star_data["r3"] = h_cnt, r_cnt
                elif "2-bintang" in label or "2-star" in label:
                    star_data["h2"], star_data["r2"] = h_cnt, r_cnt
                elif "1-bintang" in label or "1-star" in label:
                    star_data["h1"], star_data["r1"] = h_cnt, r_cnt
                elif "orkid" in label or "orchid" in label:
                    star_data["horchid"] += h_cnt
                    star_data["rorchid"] += r_cnt
                elif "unrated" in label or "tiada penarafan" in label:
                    star_data["hunrated"], star_data["runrated"] = h_cnt, r_cnt
                elif "jumlah" in label or "total" in label:
                    star_data["htotal"], star_data["rtotal"] = h_cnt, r_cnt

        tot_rooms = star_data["rtotal"] if star_data["rtotal"] > 0 else (
            star_data["r5"] + star_data["r4"] + star_data["r3"] + star_data["r2"] + star_data["r1"] + star_data["rorchid"] + star_data["runrated"]
        )
        tot_hotels = star_data["htotal"] if star_data["htotal"] > 0 else (
            star_data["h5"] + star_data["h4"] + star_data["h3"] + star_data["h2"] + star_data["h1"] + star_data["horchid"] + star_data["hunrated"]
        )

        luxury_rooms = star_data["r5"] + star_data["r4"]
        luxury_pct = round((luxury_rooms / tot_rooms) * 100.0, 2) if tot_rooms > 0 else 0.0
        midscale_pct = round((star_data["r3"] / tot_rooms) * 100.0, 2) if tot_rooms > 0 else 0.0
        budget_pct = round(100.0 - luxury_pct - midscale_pct, 2) if tot_rooms > 0 else 0.0

        hot_rec = {
            "year": year,
            "state": meta["name"],
            "state_code": meta["code"],
            "region": meta["region"],
            "hotels_5star": star_data["h5"],
            "rooms_5star": star_data["r5"],
            "hotels_4star": star_data["h4"],
            "rooms_4star": star_data["r4"],
            "hotels_3star": star_data["h3"],
            "rooms_3star": star_data["r3"],
            "hotels_2star": star_data["h2"],
            "rooms_2star": star_data["r2"],
            "hotels_1star": star_data["h1"],
            "rooms_1star": star_data["r1"],
            "hotels_orchid": star_data["horchid"],
            "rooms_orchid": star_data["rorchid"],
            "hotels_unrated": star_data["hunrated"],
            "rooms_unrated": star_data["runrated"],
            "total_hotels": tot_hotels,
            "total_rooms": tot_rooms,
            "luxury_room_share_pct": luxury_pct,
            "midscale_room_share_pct": midscale_pct,
            "budget_room_share_pct": budget_pct,
        }

    wb.close()
    return inc_rec, hot_rec


def run_pipeline():
    print("=" * 70)
    print("Parsing DTS Tourist Income (Jadual 13a) & Hotel Star Inventory (Jadual 14)")
    print("=" * 70)

    years = sorted([int(p.name) for p in DTS_BASE_DIR.glob("*") if p.is_dir() and p.name.isdigit()])
    print(f"Discovered {len(years)} DTS survey years: {years}\n")

    all_income_records = []
    all_hotel_records = []

    for y in years:
        state_dir = DTS_BASE_DIR / f"{y}/state"
        state_files = sorted(list(state_dir.glob("*.xlsx")))
        print(f"Year {y}: processing {len(state_files)} state files with ProcessPool...")

        with ProcessPoolExecutor(max_workers=4) as executor:
            futures = [executor.submit(parse_single_workbook, f, y) for f in state_files]
            inc_cnt, hot_cnt = 0, 0
            for fut in futures:
                inc_r, hot_r = fut.result()
                if inc_r:
                    all_income_records.append(inc_r)
                    inc_cnt += 1
                if hot_r:
                    all_hotel_records.append(hot_r)
                    hot_cnt += 1

        print(f"  -> Extracted income profiles: {inc_cnt}, hotel star inventories: {hot_cnt}")

    df_income = pd.DataFrame(all_income_records)
    df_hotel = pd.DataFrame(all_hotel_records)

    print("\nExtraction Summary:")
    print(f"  1. state_tourist_income_panel: {len(df_income)} records ({df_income['year'].nunique()} years, {df_income['state'].nunique()} states)")
    print(f"  2. state_hotel_star_inventory: {len(df_hotel)} records ({df_hotel['year'].nunique()} years, {df_hotel['state'].nunique()} states)")

    # Export to CSV
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    df_income.to_csv(PROCESSED_DIR / "state_tourist_income_panel.csv", index=False)
    df_hotel.to_csv(PROCESSED_DIR / "state_hotel_star_inventory.csv", index=False)

    # Insert / Replace in DuckDB and write Parquet
    con = duckdb.connect(str(DUCKDB_PATH))
    con.execute("CREATE OR REPLACE TABLE state_tourist_income_panel AS SELECT * FROM df_income")
    con.execute("CREATE OR REPLACE TABLE state_hotel_star_inventory AS SELECT * FROM df_hotel")

    con.execute(f"COPY state_tourist_income_panel TO '{PROCESSED_DIR}/state_tourist_income_panel.parquet' (FORMAT PARQUET)")
    con.execute(f"COPY state_hotel_star_inventory TO '{PROCESSED_DIR}/state_hotel_star_inventory.parquet' (FORMAT PARQUET)")

    print("\nDuckDB Tables & Parquet Files Created Successfully:")
    for tbl in ["state_tourist_income_panel", "state_hotel_star_inventory"]:
        n = con.execute(f"SELECT COUNT(*) FROM {tbl}").fetchone()[0]
        print(f"  ✓ {tbl:<35} {n:>6d} rows")
    con.close()
    print("=" * 70)


if __name__ == "__main__":
    run_pipeline()
