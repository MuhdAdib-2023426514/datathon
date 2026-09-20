"""
Excel & Parquet Inspector MCP Server.
Provides interactive inspection tools for raw Excel workbooks and Parquet datasets:
sheet enumeration, cell range viewing, header/metadata detection, and Parquet schema profiling.
"""

import sys
import io
from pathlib import Path
from typing import Optional
import pandas as pd
import openpyxl
from mcp.server.mcpserver import MCPServer

mcp = MCPServer(
    "excel_inspector",
    instructions=(
        "Specialized file inspector for raw Excel workbooks and Parquet datasets. "
        "Allows fast inspection of sheets, ranges, headers, and schemas without "
        "running heavy ingestion scripts."
    ),
)


@mcp.tool()
def list_sheets(file_path: str) -> str:
    """List all sheet names and dimensions (max row x max col) in an Excel workbook."""
    p = Path(file_path)
    if not p.exists():
        return f"Error: File not found at '{file_path}'"

    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        lines = [f"Workbook: {p.name}", f"Total Sheets: {len(wb.sheetnames)}", ""]
        for idx, name in enumerate(wb.sheetnames, 1):
            try:
                ws = wb[name]
                lines.append(f"  {idx}. {name} (approx {ws.max_row} rows x {ws.max_column} cols)")
            except Exception:
                lines.append(f"  {idx}. {name}")
        wb.close()
        return "\n".join(lines)
    except Exception as e:
        return f"Error reading Excel workbook: {e}"


@mcp.tool()
def read_range(
    file_path: str,
    sheet: str,
    start_row: int = 1,
    end_row: int = 15,
    max_cols: int = 15,
) -> str:
    """Read a specific row slice from an Excel sheet and return as formatted table.
    Row numbers are 1-indexed.
    """
    p = Path(file_path)
    if not p.exists():
        return f"Error: File not found at '{file_path}'"

    if start_row < 1:
        start_row = 1
    if end_row < start_row:
        end_row = start_row + 10
    nrows = end_row - start_row + 1

    try:
        df = pd.read_excel(
            str(p),
            sheet_name=sheet,
            header=None,
            skiprows=start_row - 1,
            nrows=nrows,
        )
        if df.shape[1] > max_cols:
            df = df.iloc[:, :max_cols]

        # Format rows with 1-based row labels matching original sheet
        df.index = range(start_row, start_row + len(df))
        return (
            f"File: {p.name} | Sheet: {sheet} | Rows {start_row} to {start_row + len(df) - 1}\n\n"
            + df.to_string()
        )
    except Exception as e:
        return f"Error reading sheet '{sheet}' from '{file_path}': {e}"


@mcp.tool()
def detect_headers(file_path: str, sheet: str, scan_rows: int = 15) -> str:
    """Inspect top N rows of an Excel sheet to identify potential header rows,
    non-empty cell counts, and candidate column labels.
    """
    p = Path(file_path)
    if not p.exists():
        return f"Error: File not found at '{file_path}'"

    try:
        df = pd.read_excel(
            str(p),
            sheet_name=sheet,
            header=None,
            nrows=scan_rows,
        )

        report = [
            f"Header Candidate Analysis for '{p.name}' -> [{sheet}]",
            "=" * 60,
        ]

        for r_idx in range(len(df)):
            row = df.iloc[r_idx]
            non_nulls = row.dropna()
            str_vals = [str(v).strip() for v in non_nulls if str(v).strip()]
            num_filled = len(str_vals)
            sample_preview = " | ".join(str_vals[:6])
            if len(str_vals) > 6:
                sample_preview += f" ... (+{len(str_vals) - 6} more)"

            flag = ""
            if num_filled >= 3:
                # Likely header or title
                if any(any(c.isalpha() for c in s) for s in str_vals):
                    flag = " <-- CANDIDATE HEADER"

            report.append(
                f"Row {r_idx + 1:2d} ({num_filled:2d} filled): {sample_preview}{flag}"
            )

        return "\n".join(report)
    except Exception as e:
        return f"Error inspecting headers: {e}"


@mcp.tool()
def read_parquet_schema(file_path: str) -> str:
    """Show column names, data types, row count, null counts, and memory footprint for a Parquet file."""
    p = Path(file_path)
    if not p.exists():
        return f"Error: File not found at '{file_path}'"

    try:
        df = pd.read_parquet(str(p))
        mem_mb = df.memory_usage(deep=True).sum() / (1024 * 1024)
        lines = [
            f"Parquet File: {p.name}",
            f"Total Rows: {len(df):,}",
            f"Total Columns: {len(df.columns)}",
            f"Memory Footprint: {mem_mb:.2f} MB",
            "",
            "Columns & Types:",
            f"{'Column':<35} {'Dtype':<15} {'Nulls':<10} {'Null %':<10}",
            "-" * 72,
        ]

        for col in df.columns:
            nulls = int(df[col].isna().sum())
            null_pct = (nulls / len(df) * 100) if len(df) > 0 else 0.0
            lines.append(
                f"{str(col):<35} {str(df[col].dtype):<15} {nulls:<10} {null_pct:>6.2f}%"
            )

        return "\n".join(lines)
    except Exception as e:
        return f"Error reading Parquet schema from '{file_path}': {e}"


@mcp.tool()
def sample_parquet(file_path: str, n_rows: int = 5) -> str:
    """Return top N rows of a Parquet file formatted as a readable table."""
    p = Path(file_path)
    if not p.exists():
        return f"Error: File not found at '{file_path}'"

    try:
        df = pd.read_parquet(str(p))
        sample = df.head(n_rows)
        return (
            f"Parquet Preview ({p.name}, showing {len(sample)} of {len(df)} rows):\n\n"
            + sample.to_string(index=False)
        )
    except Exception as e:
        return f"Error sampling Parquet file '{file_path}': {e}"


if __name__ == "__main__":
    mcp.run(transport="stdio")
